# ruff: noqa: E501
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/resumes")
BLUE = RGBColor(36, 76, 139)
DARK = RGBColor(38, 38, 38)


def set_font(run, size=9.4, bold=False, italic=False, color=DARK):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "244C8B")
    bottom.set(qn("w:space"), "3")
    borders.append(bottom)


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    set_font(r, 10.5, bold=True, color=BLUE)
    bottom_border(p)


def line(doc, left, right="", italic=False):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.1)
    table.columns[1].width = Inches(1.4)
    table._tbl.tblPr.append(OxmlElement("w:tblLayout"))
    for cell in table.rows[0].cells:
        cell.margin_top = 0
        cell.margin_bottom = 0
    p1, p2 = table.rows[0].cells[0].paragraphs[0], table.rows[0].cells[1].paragraphs[0]
    p1.paragraph_format.space_after = p2.paragraph_format.space_after = Pt(0)
    r1, r2 = p1.add_run(left), p2.add_run(right)
    set_font(r1, bold=not italic, italic=italic, color=BLUE if italic else DARK)
    set_font(r2, bold=True)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), 9.15)


def labeled(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(label + ": "), 9.15, bold=True)
    set_font(p.add_run(text), 9.15)


def build(kind):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.42)
    sec.bottom_margin = Inches(0.42)
    sec.left_margin = Inches(0.5)
    sec.right_margin = Inches(0.5)
    sec.header_distance = Inches(0.2)
    sec.footer_distance = Inches(0.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.4)
    normal.paragraph_format.space_after = Pt(0)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    set_font(name.add_run("Raymond Frias-Amaro"), 23, bold=True, color=BLUE)
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(4)
    set_font(
        contact.add_run(
            "rfrias1869@gmail.com  |  github.com/webDeveloper42  |  raymondfrias.com  |  Suitland, MD 20746"
        ),
        8.8,
        bold=True,
    )

    if kind == "it":
        title = "ENTRY-LEVEL IT SUPPORT / HELP DESK ANALYST"
        summary = (
            "Bilingual entry-level IT support professional with hands-on experience building and operating a secure help desk, "
            "administering Linux systems, troubleshooting hardware/software and connectivity issues, and documenting repeatable resolutions. "
            "Combines 4+ years of customer service with practical ticket triage, access control, MySQL, and QA experience."
        )
        skills = [
            (
                "Technical Support",
                "Tier 1 triage, incident and request tracking, hardware/software diagnostics, password resets, remote support, SLA awareness",
            ),
            (
                "Systems & Networking",
                "Windows, macOS, Ubuntu/Debian Server, Microsoft 365, printers/peripherals, DNS A/CNAME, TCP/IP basics, SSH, UFW",
            ),
            (
                "Administration & Data",
                "Linux CLI, Bash, users/groups/permissions, filesystems, MySQL, SQLite, Python virtual environments, service health checks",
            ),
            (
                "Support Operations",
                "Ticket ownership, prioritization, escalation, audit logs, SOPs, knowledge-base documentation, English/Spanish communication",
            ),
        ]
        exp_bullets = [
            "Built TicketDesk, a Python/Flask and SQLite service desk for incident intake, searchable queues, assignment, priority, SLA targets, public replies, internal notes, and resolution tracking.",
            "Implemented requester, agent, and administrator roles with object-level authorization, CSRF protection, PBKDF2 password hashing, secure headers, validation, and auditable login/ticket events.",
            "Created an idempotent one-command local setup that builds a Python virtual environment, generates secrets, initializes SQLite, seeds demo accounts and tickets, and starts the service.",
            "Designed and executed a risk-based QA session covering access control, ticket lifecycle, communication privacy, accessibility, responsive usability, persistence, and recovery; automated 20 tests at 91% coverage.",
            "Launched and operate discover-yourself.app, including DNS A/CNAME records, GitHub Actions deployment, production troubleshooting, and customer-facing support for a live product.",
        ]
        project_name = "Home Lab Server"
        project_link = "github.com/webDeveloper42/home-lab-server  |  github.com/webDeveloper42/linux-sysadmin-notes"
        project_bullets = [
            "Administered Ubuntu Server with SSH keys, UFW, users/groups, permissions, MySQL-backed services, and documented recovery procedures.",
            "Diagnosed remote-access, filesystem permission, firewall, and service configuration failures using command-line tools and logs.",
        ]
    else:
        title = "ENTRY-LEVEL SOFTWARE ENGINEER"
        summary = (
            "Entry-level software engineer with hands-on ownership of full-stack applications from design through deployment and QA. "
            "Builds secure, maintainable systems with Python, JavaScript, React, Flask, REST APIs, SQL/NoSQL databases, OOP, automated tests, and CI/CD. "
            "Bilingual: English and Spanish."
        )
        skills = [
            (
                "Languages & Frameworks",
                "Python, JavaScript, HTML5, CSS3, Flask, React, Node.js, Express, Jinja",
            ),
            (
                "Data & APIs",
                "MySQL, SQLAlchemy ORM, MongoDB, RESTful API design, third-party API integration, relational modeling",
            ),
            (
                "Engineering",
                "OOP, application factories, service layer, RBAC, BEM, responsive design, accessibility, security headers, CSRF protection",
            ),
            (
                "Testing & Delivery",
                "pytest, fixtures, integration/security testing, 91% coverage, Ruff, Bandit, Git, GitHub Actions, Python virtual environments, Render",
            ),
        ]
        exp_bullets = [
            "Architected and launched discover-yourself.app, a monetized full-stack MERN application with secure accounts, REST APIs, responsive UI, custom DNS, and GitHub Actions deployment to Render.",
            "Engineered TicketDesk, a role-based IT service management application using Python, Flask, SQLAlchemy, SQLite, Jinja, and BEM CSS with ticket intake, search, assignment, SLA tracking, comments, and audit history.",
            "Applied an OOP service layer and application-factory/blueprint architecture to separate HTTP, domain workflow, persistence, and presentation concerns.",
            "Hardened authentication and authorization with PBKDF2 password hashing, CSRF controls, parameterized ORM queries, object-level permissions, input limits, and restrictive browser security headers.",
            "Built 20 model, request-integration, workflow, and abuse-case tests reaching 91% coverage; added Ruff and Bandit quality gates plus a structured exploratory QA plan.",
            "Automated local setup with a Python virtual environment and SQLite; the idempotent script generates secrets, creates tables, seeds demo data, and launches the service.",
        ]

    section(doc, title)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(summary), 9.15)

    section(doc, "Technical Skills")
    for label, text in skills:
        labeled(doc, label, text)

    section(doc, "Professional Experience")
    role = (
        "IT Systems Developer / Software Engineer, Founder - Discover Yourself"
        if kind == "it"
        else "Software Engineer, Founder - Discover Yourself Full-Stack MERN Application"
    )
    line(doc, role, "10/2025 - Present")
    line(doc, "discover-yourself.app", italic=True)
    for item in exp_bullets:
        bullet(doc, item)

    if kind == "it":
        section(doc, "Selected Project")
        line(doc, project_name)
        line(doc, project_link, italic=True)
        for item in project_bullets:
            bullet(doc, item)

    section(doc, "Education & Leadership")
    line(doc, "Software Engineer Certificate, TripleTen", "April 2026")
    line(doc, "Eagle Scout, Boy Scouts of America - Troop 1869, Washington, DC", "2018")

    path = OUT / (
        "Raymond_Frias_Resume_IT_TicketDesk.docx"
        if kind == "it"
        else "Raymond_Frias_Resume_SE_TicketDesk.docx"
    )
    doc.save(path)
    return path


OUT.mkdir(parents=True, exist_ok=True)
requested = sys.argv[1:] or ["it", "se"]
for resume_kind in requested:
    if resume_kind not in {"it", "se"}:
        raise SystemExit("Usage: build_resumes.py [it] [se]")
    build(resume_kind)
